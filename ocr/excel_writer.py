"""Write what was extracted to a multi-sheet workbook.

Nothing here knows what the fields are called. Sheet 1's headers are the
invoice dict's own keys and Sheet 2's are the bill table's own column names,
so a form that prints different labels produces a different - and still
correct - workbook. Invoice fields that are missing are red; bill rows with
validation_flags are yellow; LR sheets where the LR number could not be read
carry a red warning at the top.
"""

import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

GREY = "F2F2F2"
BLUE = "ADD8E6"
RED = "FF4444"
YELLOW = "FFD700"

MAX_COLUMN_WIDTH = 50

# Sheet 2 opens with the metadata block, then a blank row, then the table.
METADATA_HEADERS = ("Field", "Value")
SECTION_GAP = 1


def _fill(colour: str) -> PatternFill:
    """Solid fill from a 6-digit hex, with the alpha byte Excel expects."""
    argb = colour if len(colour) == 8 else "FF" + colour
    return PatternFill(start_color=argb, end_color=argb, fill_type="solid")


def _cell_value(value):
    """Render a field for a spreadsheet cell."""
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(item) for item in value)
    return value


def make_header(key: str) -> str:
    """Turn a field key into a column title: ``lr_date`` -> ``Lr Date``."""
    return str(key).replace("_", " ").strip().title()


def _write_header(sheet, headers, row: int = 1, freeze: bool = True) -> None:
    """Bold grey header row, with everything above it frozen in place."""
    header_fill = _fill(GREY)
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=column, value=title)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    if freeze:
        sheet.freeze_panes = f"A{row + 1}"


def _autofit(sheet, start_row: int = 1) -> None:
    """Approximate auto-fit; openpyxl cannot measure rendered text."""
    for column_cells in sheet.columns:
        widths = [
            len(str(cell.value if cell.value is not None else ""))
            for cell in column_cells
            if cell.row >= start_row and hasattr(cell, 'column_letter')
        ]
        first = next((c for c in column_cells if hasattr(c, 'column_letter')), None)
        if first is None:
            continue
        sheet.column_dimensions[first.column_letter].width = min(max(widths or [0]) + 2, MAX_COLUMN_WIDTH)


def _write_invoice_sheet(sheet, invoice_data: dict, offset: int = 0) -> None:
    """The form as one wide row: its labels across the top, its values under."""
    header_row = offset + 1
    data_row = offset + 2

    if not invoice_data:
        sheet.cell(row=header_row, column=1, value="No invoice data extracted")
        return

    keys = list(invoice_data)
    _write_header(sheet, [make_header(key) for key in keys], row=header_row)

    missing = _fill(RED)
    for column, key in enumerate(keys, start=1):
        value = invoice_data[key]
        cell = sheet.cell(row=data_row, column=column, value=_cell_value(value))
        if value is None:
            cell.fill = missing

    _autofit(sheet)


def _write_bill_sheet(sheet, bill_header: dict, bill_rows: list) -> int:
    """The metadata block, a blank row, then the table. Returns the table's row."""
    _write_header(sheet, METADATA_HEADERS, freeze=False)

    metadata_fill = _fill(BLUE)
    row_number = 1
    for row_number, (key, value) in enumerate(bill_header.items(), start=2):
        label = sheet.cell(row=row_number, column=1, value=make_header(key))
        field = sheet.cell(row=row_number, column=2, value=_cell_value(value))
        label.fill = metadata_fill
        field.fill = metadata_fill

    table_row = max(row_number, 1) + SECTION_GAP + 1

    if not bill_rows:
        sheet.cell(row=table_row, column=1, value="No bill rows extracted")
        _autofit(sheet)
        return table_row

    # The columns are whatever the table named, in the order the first row
    # carries them; later rows may add a column of their own that the first
    # row did not have - nothing here needs to know any column's name.
    columns = list(bill_rows[0])
    for record in bill_rows[1:]:
        for key in record:
            if key not in columns:
                columns.append(key)

    flagged_fill = _fill(YELLOW)
    # validation_flags is internal bookkeeping — never written as a data column.
    display_columns = [c for c in columns if c != 'validation_flags']
    _write_header(sheet, [make_header(c) for c in display_columns], row=table_row)

    for offset, record in enumerate(bill_rows, start=1):
        has_flags = bool(record.get('validation_flags'))
        for column_index, column in enumerate(display_columns, start=1):
            cell = sheet.cell(row=table_row + offset, column=column_index,
                              value=_cell_value(record.get(column)))
            if has_flags:
                cell.fill = flagged_fill

    _autofit(sheet)
    return table_row


def _write_lr_sheet(sheet, lr_record: dict) -> None:
    """One Lorry Receipt's full extracted result as a vertical Field/Value
    list - every key ocr/lr_extractor.py's extract_lr produced for this
    page, in the order it produced them. Nothing here knows what those keys
    are, same as the rest of this module - a field extract_lr adds later
    still shows up here with no change needed."""
    record = lr_record or {}
    lr_no = record.get('lr_no')
    start_row = 1

    if not lr_no:
        warning_cell = sheet.cell(row=1, column=1,
                                  value="WARNING: LR number could not be read from this page. "
                                        "Fields below may be missing or misaligned.")
        warning_cell.fill = _fill(RED)
        warning_cell.font = Font(bold=True)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        start_row = 2

    _write_header(sheet, METADATA_HEADERS, row=start_row, freeze=False)
    for row_number, (key, value) in enumerate(record.items(), start=start_row + 1):
        sheet.cell(row=row_number, column=1, value=make_header(key))
        sheet.cell(row=row_number, column=2, value=_cell_value(value))
    _autofit(sheet)


# Characters Excel refuses in a sheet name, and its length cap.
_SHEET_NAME_INVALID = re.compile(r"[\\/?*\[\]:]")
MAX_SHEET_NAME_LENGTH = 31


def _lr_sheet_name(lr_no, page_index: int, taken: set) -> str:
    """A unique, Excel-legal sheet name for one LR page: ``LR_<its own LR
    number>``, or ``LR_page_<n>`` when that page's LR number itself did not
    read - never a blank or duplicate name, which openpyxl refuses outright.
    """
    digits = re.sub(r"\D", "", str(lr_no or ""))
    digits = digits.lstrip("0") or digits
    base = f"LR_{digits}" if digits else f"LR_page_{page_index + 1}"
    base = _SHEET_NAME_INVALID.sub("", base)[:MAX_SHEET_NAME_LENGTH] or f"LR_page_{page_index + 1}"

    name = base
    suffix = 2
    while name in taken:
        tag = f"_{suffix}"
        name = base[:MAX_SHEET_NAME_LENGTH - len(tag)] + tag
        suffix += 1
    taken.add(name)
    return name


def write_excel(invoice_data: dict, bill_header: dict, bill_rows: list,
                output_path: str, tax_invoice_rows: list = None,
                tax_invoice_header: dict = None, lr_records: list = None,
                layout_warnings: list = None) -> str:
    """Write the workbook and return ``output_path``.

    Args:
        invoice_data: Flat label/value dict from the cover form.
        bill_header: Flat label/value dict from above the bill table.
        bill_rows: One dict per table row, keyed by the table's own columns.
        output_path: Where to save. Parent directories are created.
        tax_invoice_rows: One dict per row of the Transportation Tax
            Invoice table (see ocr/tax_invoice_extractor.py), keyed by its
            own columns. A third "Tax Invoice" sheet is added only when
            this is non-empty - a combined PDF with no such page (or one
            not yet run through that extractor) still gets the same
            two-sheet workbook as before.
        tax_invoice_header: Flat label/value dict for that sheet's own
            metadata block, if any - empty by default, since that table
            carries no such block above it the way the bill table does.
        lr_records: One dict per individual Lorry Receipt page (see
            ocr/lr_extractor.py's extract_lr), in page order. Each gets its
            own sheet, named after its own ``lr_no`` field - not merged
            into one table, since every LR is its own document with its
            own metadata block and handwritten acknowledgement box, not a
            repeating row of a shared table the way the bill/tax invoice
            are.
    """
    workbook = Workbook()

    invoice_sheet = workbook.active
    invoice_sheet.title = "Invoice"

    if layout_warnings:
        for i, warning in enumerate(layout_warnings, start=1):
            cell = invoice_sheet.cell(row=i, column=1, value=f"WARNING: {warning}")
            cell.fill = _fill(RED)
            cell.font = Font(bold=True)

    _write_invoice_sheet(invoice_sheet, invoice_data or {}, offset=len(layout_warnings or []))

    _write_bill_sheet(workbook.create_sheet("Bill"), bill_header or {}, bill_rows or [])

    if tax_invoice_rows:
        _write_bill_sheet(workbook.create_sheet("Tax Invoice"),
                          tax_invoice_header or {}, tax_invoice_rows)

    if lr_records:
        taken_names = set(workbook.sheetnames)
        for index, record in enumerate(lr_records):
            sheet_name = _lr_sheet_name(record.get("lr_no"), index, taken_names)
            _write_lr_sheet(workbook.create_sheet(sheet_name), record)

    directory = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(directory, exist_ok=True)
    workbook.save(output_path)
    return output_path


if __name__ == "__main__":
    import tempfile

    from openpyxl import load_workbook

    failures = []

    def check(label, actual, expected):
        ok = actual == expected
        if not ok:
            failures.append(label)
        print(f"[{'ok  ' if ok else 'FAIL'}] {label}: {actual!r}"
              + ("" if ok else f" != {expected!r}"))

    def rgb(cell):
        colour = cell.fill.start_color.rgb
        return colour[2:] if isinstance(colour, str) and len(colour) == 8 else colour

    invoice_data = {
        "Vendor_Name": "INLAND WORLD LOGISTICS P LTD.",
        "Vendor_Code": "20050923",
        "Invoice_Date": None,
        "PO_Type_Selected": "PO6X",
        "Stamp_Jsw_Steel_Ltd": "JSW STEEL LTD RAIGARH 26 JUN 2026",
    }
    bill_header = {
        "Bill_No": "INRGR2600002",
        "Bill_Date": "10/06/2026",
        "column_numbers": ["1", "2", "3"],
    }
    bill_rows = [
        {"sr_no": "1", "shipment_no": "0000001350", "amount": "84008.60"},
        {"sr_no": "2", "shipment_no": None, "shipment_no_raw": "OOOO135l",
         "amount": "69448.20"},
    ]

    output_path = os.path.join(tempfile.mkdtemp(), "nested", "report.xlsx")
    write_excel(invoice_data, bill_header, bill_rows, output_path)
    workbook = load_workbook(output_path)

    print("--- workbook ---")
    check("file written", os.path.isfile(output_path), True)
    check("parent dirs created", os.path.isdir(os.path.dirname(output_path)), True)
    check("sheet names and order", workbook.sheetnames, ["Invoice", "Bill"])

    invoice_sheet = workbook["Invoice"]
    bill_sheet = workbook["Bill"]

    print("\n--- Sheet 1: Invoice ---")
    check("headers are the dict's own keys", [c.value for c in invoice_sheet[1]],
          [make_header(key) for key in invoice_data])
    check("header bold", invoice_sheet.cell(row=1, column=1).font.bold, True)
    check("header grey", rgb(invoice_sheet.cell(row=1, column=1)), GREY)
    check("frozen below header", invoice_sheet.freeze_panes, "A2")
    check("one value row", invoice_sheet.max_row, 2)
    check("value written", invoice_sheet.cell(row=2, column=1).value,
          "INLAND WORLD LOGISTICS P LTD.")
    check("None cell RED", rgb(invoice_sheet.cell(row=2, column=3)), RED)
    check("present cell unfilled", invoice_sheet.cell(row=2, column=1).fill.fill_type, None)

    print("\n--- Sheet 2: Bill ---")
    table_row = len(bill_header) + 1 + SECTION_GAP + 1
    check("metadata headers", [bill_sheet.cell(row=1, column=c).value for c in (1, 2)],
          list(METADATA_HEADERS))
    check("metadata label", bill_sheet.cell(row=2, column=1).value, "Bill No")
    check("metadata value", bill_sheet.cell(row=2, column=2).value, "INRGR2600002")
    check("list rendered as text", bill_sheet.cell(row=4, column=2).value, "1, 2, 3")
    check("metadata BLUE", rgb(bill_sheet.cell(row=2, column=2)), BLUE)
    check("blank row before table",
          bill_sheet.cell(row=table_row - 1, column=1).value, None)
    check("table headers from the rows' own columns",
          [c.value for c in bill_sheet[table_row]],
          ["Sr No", "Shipment No", "Amount", "Shipment No Raw"])
    check("table header bold", bill_sheet.cell(row=table_row, column=1).font.bold, True)
    check("frozen below table header", bill_sheet.freeze_panes, f"A{table_row + 1}")
    check("data row", bill_sheet.cell(row=table_row + 1, column=2).value, "0000001350")
    check("no conditional fill on a failed cell",
          bill_sheet.cell(row=table_row + 2, column=2).fill.fill_type, None)
    check("raw kept beside it", bill_sheet.cell(row=table_row + 2, column=4).value,
          "OOOO135l")
    check("row count", bill_sheet.max_row, table_row + len(bill_rows))

    print("\n--- Sheet(s) 3+: one per LR ---")
    lr_records = [
        {"lr_no": "0000001350", "truck_no": "OD02CY4777", "net_wt": "41.08"},
        {"lr_no": None, "truck_no": "CG04QC9560", "net_wt": "33.96"},
        {"lr_no": "0000001350", "truck_no": "DUPLICATE", "net_wt": "0"},
    ]
    lr_path = os.path.join(tempfile.mkdtemp(), "lr.xlsx")
    write_excel(invoice_data, bill_header, bill_rows, lr_path, lr_records=lr_records)
    lr_workbook = load_workbook(lr_path)
    check("one sheet per LR, named by its own LR No",
          lr_workbook.sheetnames, ["Invoice", "Bill", "LR_1350", "LR_page_2", "LR_1350_2"])
    lr_sheet = lr_workbook["LR_1350"]
    check("LR sheet headers", [lr_sheet.cell(row=1, column=c).value for c in (1, 2)],
          list(METADATA_HEADERS))
    check("LR field written", lr_sheet.cell(row=2, column=1).value, "Lr No")
    check("LR value written", lr_sheet.cell(row=2, column=2).value, "0000001350")
    check("second LR field", lr_sheet.cell(row=3, column=1).value, "Truck No")

    print("\n--- empty input ---")
    empty_path = os.path.join(tempfile.mkdtemp(), "empty.xlsx")
    write_excel({}, {}, [], empty_path)
    empty = load_workbook(empty_path)
    check("sheets still created", empty.sheetnames, ["Invoice", "Bill"])
    check("invoice sheet says so", empty["Invoice"].cell(row=1, column=1).value,
          "No invoice data extracted")
    check("bill sheet says so", empty["Bill"].cell(row=3, column=1).value,
          "No bill rows extracted")

    print()
    print(f"workbook: {output_path}")
    print(f"{len(failures)} FAILED: {failures}" if failures else "all checks passed")
